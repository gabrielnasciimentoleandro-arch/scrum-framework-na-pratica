#!/usr/bin/env python3
"""Gera as versões DOCX e PDF do guia Framework Scrum na Prática."""

from __future__ import annotations

import html
import os
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from importlib.util import find_spec
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GUIA_COMPLETO.md"
OUTPUT_DIR = ROOT / "entrega"
COVER_PNG = ROOT / "assets" / "capa-documento.png"
DOCX_PATH = OUTPUT_DIR / "Guia_Framework_Scrum.docx"
PDF_PATH = OUTPUT_DIR / "Guia_Framework_Scrum.pdf"

NAVY = "22263F"
DARK = "17152F"
PURPLE = "6B4EFF"
PURPLE_DARK = "46358D"
TEAL = "13A085"
CORAL = "F06449"
GOLD = "D08A18"
MUTED = "66708B"
LIGHT = "F4F5FB"
WHITE = "FFFFFF"


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    number: str = ""
    rows: list[list[str]] = field(default_factory=list)


INLINE_RE = re.compile(
    r"\*\*(.+?)\*\*|\*([^*]+?)\*|\[([^\]]+)\]\((https?://[^)]+)\)|"
    r"<(https?://[^>]+)>|`([^`]+)`|\[\^(\d+)\]"
)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells
    )


def parse_markdown(path: Path) -> list[Block]:
    """Interpreta o subconjunto de Markdown usado no guia."""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    paragraph: list[str] = []
    started = False
    index = 0

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            blocks.append(
                Block("paragraph", " ".join(item.strip() for item in paragraph))
            )
            paragraph = []

    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("## Resumo"):
            started = True
        if not started:
            index += 1
            continue

        if not line.strip():
            flush()
            index += 1
            continue

        if line.lstrip().startswith("|"):
            flush()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index].rstrip())
                index += 1
            parsed = [parse_table_row(item) for item in table_lines]
            rows = [row for row in parsed if not is_table_separator(row)]
            if rows:
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                blocks.append(Block("table", rows=rows))
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush()
            blocks.append(
                Block("heading", heading.group(2).strip(), len(heading.group(1)) - 1)
            )
            index += 1
            continue

        if line.strip() == "---":
            flush()
            blocks.append(Block("rule"))
            index += 1
            continue

        if line.startswith(">"):
            flush()
            quote: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith(">"):
                quote.append(lines[index].lstrip()[1:].strip())
                index += 1
            blocks.append(Block("quote", " ".join(quote)))
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if bullet:
            flush()
            blocks.append(Block("bullet", bullet.group(1)))
            index += 1
            continue

        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if numbered:
            flush()
            blocks.append(Block("number", numbered.group(2), number=numbered.group(1)))
            index += 1
            continue

        paragraph.append(line)
        index += 1

    flush()
    return blocks


def find_font(*names: str) -> Path:
    """Localiza fonte em ReportLab, Linux, macOS ou Windows."""
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts/truetype/liberation2"),
        Path("/usr/share/fonts/truetype/liberation"),
        Path("/usr/local/share/fonts"),
        Path.home() / ".fonts",
        Path.home() / "Library" / "Fonts",
        Path("/Library/Fonts"),
        Path("/System/Library/Fonts"),
        Path("/System/Library/Fonts/Supplemental"),
    ]
    windows_dir = os.environ.get("WINDIR")
    candidates.append(
        Path(windows_dir) / "Fonts" if windows_dir else Path("C:/Windows/Fonts")
    )

    reportlab_spec = find_spec("reportlab")
    if reportlab_spec and reportlab_spec.submodule_search_locations:
        reportlab_root = Path(next(iter(reportlab_spec.submodule_search_locations)))
        candidates.insert(0, reportlab_root / "fonts")

    for folder in candidates:
        for name in names:
            candidate = folder / name
            if candidate.is_file():
                return candidate
        if folder.is_dir():
            available = {
                child.name.casefold(): child
                for child in folder.iterdir()
                if child.is_file()
            }
            for name in names:
                if match := available.get(name.casefold()):
                    return match
    raise FileNotFoundError(f"Fonte não encontrada: {', '.join(names)}")


def project_fonts() -> tuple[Path, Path, Path, Path]:
    regular = find_font("Vera.ttf", "DejaVuSans.ttf", "Arial.ttf", "arial.ttf")
    bold = find_font(
        "VeraBd.ttf", "DejaVuSans-Bold.ttf", "Arial Bold.ttf", "arialbd.ttf"
    )
    try:
        italic = find_font("VeraIt.ttf", "DejaVuSans-Oblique.ttf", "ariali.ttf")
    except FileNotFoundError:
        italic = regular
    try:
        bold_italic = find_font(
            "VeraBI.ttf", "DejaVuSans-BoldOblique.ttf", "arialbi.ttf"
        )
    except FileNotFoundError:
        bold_italic = bold
    return regular, bold, italic, bold_italic


def generate_cover_png(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1240, 1754
    image = Image.new("RGB", (width, height), "#111329")
    draw = ImageDraw.Draw(image)
    top, bottom = (17, 19, 41), (41, 35, 84)
    for y in range(height):
        ratio = y / (height - 1)
        color = tuple(round(top[i] * (1 - ratio) + bottom[i] * ratio) for i in range(3))
        draw.line((0, y, width, y), fill=color)

    draw.ellipse((800, -170, 1450, 480), fill="#30276a")
    draw.ellipse((850, 1160, 1500, 1810), fill="#15172f")
    draw.rounded_rectangle((85, 105, 408, 170), radius=32, fill="#13a085")

    regular, bold, _, _ = project_fonts()
    tag = ImageFont.truetype(str(bold), 25)
    kicker = ImageFont.truetype(str(bold), 27)
    title = ImageFont.truetype(str(bold), 76)
    subtitle = ImageFont.truetype(str(regular), 48)
    description = ImageFont.truetype(str(regular), 28)
    small = ImageFont.truetype(str(regular), 22)
    badge = ImageFont.truetype(str(bold), 50)

    draw.text((246, 138), "DESAFIO DE PROJETO", font=tag, fill="white", anchor="mm")
    draw.text((85, 288), "SCRUM GUIDE 2020", font=kicker, fill="#f4c95d")
    draw.text((85, 390), "Framework Scrum", font=title, fill="white")
    draw.text((85, 493), "na prática", font=subtitle, fill="#d5d2ef")
    draw.rounded_rectangle((85, 582, 760, 590), radius=4, fill="#f06449")
    draw.text(
        (85, 650),
        "Responsabilidades, eventos, artefatos,",
        font=description,
        fill="#b9b8d3",
    )
    draw.text(
        (85, 697),
        "pilares, valores e aplicação prática",
        font=description,
        fill="#b9b8d3",
    )

    cx, cy = 910, 920
    draw.ellipse((cx - 190, cy - 190, cx + 190, cy + 190), fill="#f7f7ff")
    draw.arc(
        (cx - 140, cy - 140, cx + 140, cy + 140), -80, 30, fill="#6b4eff", width=35
    )
    draw.arc(
        (cx - 140, cy - 140, cx + 140, cy + 140), 40, 150, fill="#13a085", width=35
    )
    draw.arc(
        (cx - 140, cy - 140, cx + 140, cy + 140), 160, 270, fill="#f06449", width=35
    )
    draw.text((cx, cy - 12), "ENTREGAR", font=tag, fill="#22263f", anchor="mm")
    value_font = ImageFont.truetype(str(bold), 38)
    draw.text((cx, cy + 34), "VALOR", font=value_font, fill="#6b4eff", anchor="mm")

    draw.text((85, 1515), "GUIA VISUAL AUTORAL", font=badge, fill="white")
    draw.text((85, 1620), "DIO  •  PORTFÓLIO  •  2026", font=small, fill="#8e90b2")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, optimize=True)


def iter_inline(text: str) -> Iterable[tuple[str, str, str | None]]:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            yield "text", text[cursor : match.start()], None
        if match.group(1) is not None:
            yield "bold", match.group(1), None
        elif match.group(2) is not None:
            yield "italic", match.group(2), None
        elif match.group(3) is not None:
            yield "link", match.group(3), match.group(4)
        elif match.group(5) is not None:
            yield "link", match.group(5), match.group(5)
        elif match.group(6) is not None:
            yield "code", match.group(6), None
        else:
            yield "footnote", f"[{match.group(7)}]", None
        cursor = match.end()
    if cursor < len(text):
        yield "text", text[cursor:], None


def add_docx_hyperlink(paragraph, text: str, url: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), PURPLE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_docx_inline(paragraph, text: str) -> None:
    from docx.shared import Pt, RGBColor

    for kind, value, url in iter_inline(text):
        if kind == "link":
            add_docx_hyperlink(paragraph, value, url or value)
            continue
        run = paragraph.add_run(value)
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True
        elif kind == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(PURPLE_DARK)
        elif kind == "footnote":
            run.font.superscript = True
            run.font.size = Pt(8)


def set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(
    cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def create_docx(blocks: list[Block], output: Path, cover: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    properties = document.core_properties
    properties.title = "Framework Scrum na Prática"
    properties.subject = "Desafio DIO — Completando o Framework Scrum"
    properties.author = "0Barone"
    properties.keywords = "Scrum, DIO, Sprint, Product Backlog, empirismo"
    properties.comments = "Guia autoral baseado no Scrum Guide 2020."

    first = document.sections[0]
    first.page_width, first.page_height = Cm(21), Cm(29.7)
    first.top_margin = first.bottom_margin = first.left_margin = first.right_margin = (
        Cm(0.5)
    )
    cover_paragraph = document.add_paragraph()
    cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_paragraph.paragraph_format.space_after = Pt(0)
    cover_paragraph.paragraph_format.line_spacing = Pt(1)
    cover_paragraph.add_run().add_picture(str(cover), width=Cm(20), height=Cm(28.3))

    content = document.add_section(WD_SECTION.NEW_PAGE)
    content.page_width, content.page_height = Cm(21), Cm(29.7)
    content.top_margin, content.bottom_margin = Cm(2.2), Cm(2)
    content.left_margin, content.right_margin = Cm(2.4), Cm(2.2)
    content.header_distance, content.footer_distance = Cm(0.9), Cm(0.9)
    content.header.is_linked_to_previous = False
    content.footer.is_linked_to_previous = False
    page_number_type = OxmlElement("w:pgNumType")
    page_number_type.set(qn("w:start"), "1")
    content._sectPr.append(page_number_type)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6.5)

    for style_name, size, color in (
        ("Title", 25, DARK),
        ("Heading 1", 19, DARK),
        ("Heading 2", 14.5, PURPLE_DARK),
        ("Heading 3", 11.5, TEAL),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(13 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)

    quote_style = document.styles.add_style(
        "Citação de destaque", WD_STYLE_TYPE.PARAGRAPH
    )
    quote_style.base_style = normal
    quote_style.font.size = Pt(11)
    quote_style.font.italic = True
    quote_style.font.color.rgb = RGBColor.from_string(PURPLE_DARK)
    quote_style.paragraph_format.left_indent = Cm(0.8)
    quote_style.paragraph_format.right_indent = Cm(0.4)
    quote_style.paragraph_format.space_before = (
        quote_style.paragraph_format.space_after
    ) = Pt(9)

    header = content.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("FRAMEWORK SCRUM NA PRÁTICA  •  SCRUM GUIDE 2020")
    header_run.font.name, header_run.font.size, header_run.font.bold = (
        "Aptos",
        Pt(8),
        True,
    )
    header_run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = content.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("DIO  •  PORTFÓLIO  •  ")
    footer_run.font.name, footer_run.font.size = "Aptos", Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(footer)

    title = document.add_paragraph(style="Title")
    title.add_run("Framework Scrum na Prática")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run(
        "Guia visual e aplicado • Desafio DIO • Outubro de 2026"
    )
    subtitle_run.font.size, subtitle_run.font.bold = Pt(10), True
    subtitle_run.font.color.rgb = RGBColor.from_string(CORAL)
    bar = document.add_table(rows=1, cols=1)
    set_cell_shading(bar.cell(0, 0), CORAL)
    bar.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(0)

    for block in blocks:
        if block.kind == "heading":
            style = (
                "Heading 1"
                if block.level == 1
                else "Heading 2"
                if block.level == 2
                else "Heading 3"
            )
            paragraph = document.add_paragraph(style=style)
            add_docx_inline(paragraph, block.text)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph(style="Normal")
            add_docx_inline(paragraph, block.text)
        elif block.kind == "quote":
            paragraph = document.add_paragraph(style="Citação de destaque")
            add_docx_inline(paragraph, block.text)
            props = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "20")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), PURPLE)
            borders.append(left)
            props.append(borders)
        elif block.kind in {"bullet", "number"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(-0.55)
            paragraph.paragraph_format.space_after = Pt(4)
            prefix = "• " if block.kind == "bullet" else f"{block.number}. "
            lead = paragraph.add_run(prefix)
            lead.bold = True
            lead.font.color.rgb = RGBColor.from_string(
                PURPLE if block.kind == "number" else TEAL
            )
            add_docx_inline(paragraph, block.text)
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row in enumerate(block.rows):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    if row_index == 0:
                        set_cell_shading(cell, PURPLE_DARK)
                    elif row_index % 2 == 0:
                        set_cell_shading(cell, "F0F1F8")
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_docx_inline(paragraph, value)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor.from_string(WHITE)
                if row_index == 0:
                    set_repeat_header(table.rows[0])
            after = document.add_paragraph()
            after.paragraph_format.space_after = Pt(1)
        elif block.kind == "rule":
            paragraph = document.add_paragraph()
            props = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:color"), "D9DCE9")
            borders.append(bottom)
            props.append(borders)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def rl_inline(text: str) -> str:
    chunks: list[str] = []
    for kind, value, url in iter_inline(text):
        safe = html.escape(value, quote=True)
        if kind == "bold":
            chunks.append(f"<b>{safe}</b>")
        elif kind == "italic":
            chunks.append(f"<i>{safe}</i>")
        elif kind == "link":
            safe_url = html.escape(url or value, quote=True)
            chunks.append(f'<a href="{safe_url}" color="#{PURPLE}">{safe}</a>')
        elif kind == "code":
            chunks.append(
                f'<font name="ProjectMono" color="#{PURPLE_DARK}">{safe}</font>'
            )
        elif kind == "footnote":
            chunks.append(f"<super>{safe}</super>")
        else:
            chunks.append(safe)
    return "".join(chunks)


def create_pdf(blocks: list[Block], output: Path) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            HRFlowable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise SystemExit(
            "Instale as dependências: python -m pip install -r requirements.txt"
        ) from exc

    regular, bold, italic, bold_italic = project_fonts()
    mono = find_font("VeraMono.ttf", "DejaVuSansMono.ttf", "consola.ttf")
    for name, path in (
        ("ProjectSans", regular),
        ("ProjectSans-Bold", bold),
        ("ProjectSans-Italic", italic),
        ("ProjectSans-BoldItalic", bold_italic),
        ("ProjectMono", mono),
    ):
        pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily(
        "ProjectSans",
        normal="ProjectSans",
        bold="ProjectSans-Bold",
        italic="ProjectSans-Italic",
        boldItalic="ProjectSans-BoldItalic",
    )

    page_width, page_height = A4
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=2.25 * cm,
        rightMargin=2.1 * cm,
        topMargin=2.2 * cm,
        bottomMargin=2 * cm,
        title="Framework Scrum na Prática",
        author="0Barone",
        subject="Desafio DIO — Completando o Framework Scrum",
        keywords="Scrum, DIO, Sprint, Product Backlog, empirismo",
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="ProjectSans",
        fontSize=9.8,
        leading=14.2,
        textColor=colors.HexColor(f"#{NAVY}"),
        alignment=TA_JUSTIFY,
        spaceAfter=6.5,
        splitLongWords=True,
    )
    h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontName="ProjectSans-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor(f"#{DARK}"),
        spaceBefore=14,
        spaceAfter=7,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "H2",
        parent=styles["Heading2"],
        fontName="ProjectSans-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor(f"#{PURPLE_DARK}"),
        spaceBefore=11,
        spaceAfter=5,
        keepWithNext=True,
    )
    h3 = ParagraphStyle(
        "H3",
        parent=styles["Heading3"],
        fontName="ProjectSans-Bold",
        fontSize=10.8,
        leading=14,
        textColor=colors.HexColor(f"#{TEAL}"),
        spaceBefore=9,
        spaceAfter=4,
        keepWithNext=True,
    )
    bullet = ParagraphStyle(
        "Bullet", parent=body, leftIndent=15, firstLineIndent=-11, spaceAfter=3.5
    )
    quote = ParagraphStyle(
        "Quote",
        parent=body,
        fontName="ProjectSans-Italic",
        fontSize=10.5,
        leading=15,
        leftIndent=16,
        rightIndent=10,
        borderColor=colors.HexColor(f"#{PURPLE}"),
        borderWidth=0,
        borderPadding=8,
        backColor=colors.HexColor("#F1EFFF"),
        textColor=colors.HexColor(f"#{PURPLE_DARK}"),
        spaceBefore=7,
        spaceAfter=9,
    )
    title_style = ParagraphStyle(
        "Title", parent=h1, fontSize=23, leading=28, spaceAfter=6
    )
    meta_style = ParagraphStyle(
        "Meta",
        parent=body,
        fontName="ProjectSans-Bold",
        fontSize=9,
        textColor=colors.HexColor(f"#{CORAL}"),
        alignment=TA_LEFT,
        spaceAfter=13,
    )
    table_header = ParagraphStyle(
        "TableHeader",
        parent=body,
        fontName="ProjectSans-Bold",
        fontSize=7.4,
        leading=9.4,
        textColor=colors.white,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    table_cell = ParagraphStyle(
        "TableCell",
        parent=body,
        fontSize=7.4,
        leading=9.6,
        alignment=TA_LEFT,
        spaceAfter=0,
        splitLongWords=True,
        wordWrap="CJK",
    )

    def draw_cover(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setTitle("Framework Scrum na Prática")
        canvas.setAuthor("0Barone")
        canvas.setFillColor(colors.HexColor("#111329"))
        canvas.rect(0, 0, page_width, page_height, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#30276A"))
        canvas.circle(
            page_width + 1.1 * cm, page_height - 2.5 * cm, 6 * cm, stroke=0, fill=1
        )
        canvas.setFillColor(colors.HexColor("#15172F"))
        canvas.circle(page_width + 1.4 * cm, 2.4 * cm, 6.6 * cm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(f"#{TEAL}"))
        canvas.roundRect(
            2 * cm,
            page_height - 3.6 * cm,
            5.6 * cm,
            0.75 * cm,
            0.35 * cm,
            stroke=0,
            fill=1,
        )
        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 10)
        canvas.drawCentredString(
            4.8 * cm, page_height - 3.35 * cm, "DESAFIO DE PROJETO"
        )
        canvas.setFillColor(colors.HexColor("#F4C95D"))
        canvas.setFont("ProjectSans-Bold", 11)
        canvas.drawString(2 * cm, page_height - 5.9 * cm, "SCRUM GUIDE 2020")
        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 29)
        canvas.drawString(2 * cm, page_height - 8.1 * cm, "Framework Scrum")
        canvas.setFillColor(colors.HexColor("#D5D2EF"))
        canvas.setFont("ProjectSans", 20)
        canvas.drawString(2 * cm, page_height - 9.55 * cm, "na prática")
        canvas.setFillColor(colors.HexColor(f"#{CORAL}"))
        canvas.roundRect(
            2 * cm,
            page_height - 10.65 * cm,
            12 * cm,
            0.08 * cm,
            0.04 * cm,
            stroke=0,
            fill=1,
        )
        canvas.setFillColor(colors.HexColor("#B9B8D3"))
        canvas.setFont("ProjectSans", 11.5)
        canvas.drawString(
            2 * cm, page_height - 12.1 * cm, "Responsabilidades, eventos, artefatos,"
        )
        canvas.drawString(
            2 * cm, page_height - 12.75 * cm, "pilares, valores e aplicação prática"
        )

        cx, cy = 15.7 * cm, 10.5 * cm
        canvas.setFillColor(colors.HexColor("#F7F7FF"))
        canvas.circle(cx, cy, 2.7 * cm, stroke=0, fill=1)
        canvas.setStrokeColor(colors.HexColor(f"#{PURPLE}"))
        canvas.setLineWidth(9)
        canvas.arc(cx - 2 * cm, cy - 2 * cm, cx + 2 * cm, cy + 2 * cm, -80, 100)
        canvas.setStrokeColor(colors.HexColor(f"#{TEAL}"))
        canvas.arc(cx - 2 * cm, cy - 2 * cm, cx + 2 * cm, cy + 2 * cm, 40, 100)
        canvas.setStrokeColor(colors.HexColor(f"#{CORAL}"))
        canvas.arc(cx - 2 * cm, cy - 2 * cm, cx + 2 * cm, cy + 2 * cm, 160, 100)
        canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
        canvas.setFont("ProjectSans-Bold", 8)
        canvas.drawCentredString(cx, cy + 0.15 * cm, "ENTREGAR")
        canvas.setFillColor(colors.HexColor(f"#{PURPLE}"))
        canvas.setFont("ProjectSans-Bold", 15)
        canvas.drawCentredString(cx, cy - 0.55 * cm, "VALOR")

        canvas.setFillColor(colors.white)
        canvas.setFont("ProjectSans-Bold", 12)
        canvas.drawString(2 * cm, 2.8 * cm, "GUIA VISUAL AUTORAL")
        canvas.setFillColor(colors.HexColor("#8E90B2"))
        canvas.setFont("ProjectSans", 9)
        canvas.drawString(2 * cm, 1.8 * cm, "DIO  •  PORTFÓLIO  •  2026")
        canvas.restoreState()

    def draw_content(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D9DCE9"))
        canvas.setLineWidth(0.5)
        canvas.line(
            2.25 * cm,
            page_height - 1.45 * cm,
            page_width - 2.1 * cm,
            page_height - 1.45 * cm,
        )
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.setFont("ProjectSans-Bold", 7.5)
        canvas.drawString(
            2.25 * cm,
            page_height - 1.18 * cm,
            "FRAMEWORK SCRUM NA PRÁTICA  •  SCRUM GUIDE 2020",
        )
        canvas.setFont("ProjectSans", 8)
        canvas.drawCentredString(
            page_width / 2,
            1.08 * cm,
            f"DIO  •  PORTFÓLIO  •  {canvas.getPageNumber() - 1}",
        )
        canvas.restoreState()

    story = [Spacer(1, 1), PageBreak()]
    story.append(Paragraph("Framework Scrum na Prática", title_style))
    story.append(
        Paragraph("Guia visual e aplicado • Desafio DIO • Outubro de 2026", meta_style)
    )
    story.append(
        HRFlowable(
            width="100%",
            thickness=2.2,
            color=colors.HexColor(f"#{CORAL}"),
            spaceAfter=11,
        )
    )

    for block in blocks:
        if block.kind == "heading":
            style = h1 if block.level == 1 else h2 if block.level == 2 else h3
            story.append(Paragraph(rl_inline(block.text), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(rl_inline(block.text), body))
        elif block.kind == "quote":
            story.append(Paragraph(rl_inline(block.text), quote))
        elif block.kind == "bullet":
            story.append(Paragraph(f"•&nbsp;&nbsp;{rl_inline(block.text)}", bullet))
        elif block.kind == "number":
            story.append(
                Paragraph(
                    f"<b>{block.number}.</b>&nbsp;&nbsp;{rl_inline(block.text)}", bullet
                )
            )
        elif block.kind == "rule":
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.7,
                    color=colors.HexColor("#D9DCE9"),
                    spaceBefore=7,
                    spaceAfter=7,
                )
            )
        elif block.kind == "table":
            columns = len(block.rows[0])
            usable = doc.width
            if columns == 4:
                widths = [usable * 0.1, usable * 0.35, usable * 0.39, usable * 0.16]
            elif columns == 3:
                widths = [usable * 0.24, usable * 0.5, usable * 0.26]
            elif columns == 2:
                widths = [usable * 0.38, usable * 0.62]
            else:
                widths = [usable / columns] * columns
            data = [
                [
                    Paragraph(
                        rl_inline(cell), table_header if row_index == 0 else table_cell
                    )
                    for cell in row
                ]
                for row_index, row in enumerate(block.rows)
            ]
            table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor(f"#{PURPLE_DARK}"),
                        ),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F2F3F9")],
                        ),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9DCE9")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([Spacer(1, 4), table, Spacer(1, 8)])

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story, onFirstPage=draw_cover, onLaterPages=draw_content)


def main() -> None:
    blocks = parse_markdown(SOURCE)
    if not blocks or blocks[0].text != "Resumo":
        raise SystemExit("Estrutura esperada do guia não encontrada.")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    generate_cover_png(COVER_PNG)
    create_docx(blocks, DOCX_PATH, COVER_PNG)
    create_pdf(blocks, PDF_PATH)
    print(f"Gerado: {DOCX_PATH.relative_to(ROOT)}")
    print(f"Gerado: {PDF_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
